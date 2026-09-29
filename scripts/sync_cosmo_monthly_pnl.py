#!/usr/bin/env python3
"""COSMO 월 손익(Comparative PnL 시트 1장)을 cosmo_2026.json monthly 에 넣는다.

월 손익은 단일 시트 파일(「코스모 8월 손익.xlsx」)로 오기도 하고, 누적 통합본
(「코스모 2026년 월별 손익 자료.xlsx」의 「Comparative PnL_7월」)으로 오기도 한다.
필드 규칙은 외부 생성기(cosmo-dashboard/scripts/extract_cosmo.py parse_monthly)와 같다 —
1~7월이 그 생성기로 만들어졌으므로 같은 시트를 넣으면 같은 레코드가 나와야 한다.

  python3 scripts/sync_cosmo_monthly_pnl.py "<xlsx>" --dry-run
  python3 scripts/sync_cosmo_monthly_pnl.py "<xlsx>" --sheet "Comparative PnL_7월"
  python3 scripts/sync_cosmo_monthly_pnl.py "<8월 xlsx>" \
      --prev-workbook "<월별 손익 자료.xlsx>" --prev-sheet "Comparative PnL_7월"

--prev-* 를 주면 행·부문마다 「당월 = 이번 YTD − 전월 YTD」를 대조하고, 어긋난 칸은 YTD 차분으로
복원한 뒤 인쇄값을 correction.printed 에 남긴다. 2026-08 시트는 원장 PnL-Cannery 가 Media·Ingredients 를
2행씩으로 늘렸는데 당월 열만 옛 1:1 행 참조로 남아 Can~Others 가 밀리고 캐너리 Others $4,504 가 빠졌다.
YTD 열은 새 배치를 따라 원장 당월값(외부 링크 캐시)과 센트 단위로 일치했다.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path

import openpyxl

MONTHS = ['JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE', 'JULY',
          'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER']

# 계정 라벨은 시트마다 표기가 달라(Revenue/Revenues, Operating Profit/Profits) 정규화해 매칭한다.
PNL_ROWS = {
    'revenues': 'revenue', 'revenue': 'revenue', 'costofsales': 'cos', 'grossprofit': 'gp',
    's&aexpense': 'sga', 'operatingprofits': 'op', 'operatingprofit': 'op',
    'nonoperationprofits': 'nonop', 'interestexpense': 'interest', 'netincome': 'net',
}
COST_LINES = [
    'Fish', 'Media', 'Ingredients', 'Can', 'End', 'Flat Pouch', 'Others',
    'Salary (DL)', 'Benefits (DL)', 'Salary (IDL)', 'Benefits (IDL)',
    'Electricity', 'RFO', 'Diesel', 'Water', 'Consumables', 'Repair/Maintenance',
    'Depreciation',
]
# 매출 세부 계정 — Fish Waste 는 원가 구간에도 같은 라벨이 있어 첫 등장(매출 구간)만 쓴다
REVENUE_LINES = [
    'Export Sales', 'Local Sales', 'Precooked Loin Sales', 'Pouch', 'Fish Oil', 'Fish Waste',
    'Fish Heads', 'Raw Fish Sales', 'Sales Returns', 'Sales Allowances',
]
UNITS = ('Cannery', 'Fishmeal', 'CBU', 'FBU', 'Total')


def norm(v) -> str:
    return '' if v is None else re.sub(r'\s+', '', unicodedata.normalize('NFC', str(v)))


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    s = str(v).strip().replace(',', '')
    try:
        f = float(s)
    except ValueError:
        return None
    return f if math.isfinite(f) else None


def rnd(v, d=2):
    return None if v is None else round(v, d)


def title_month(ws) -> tuple[int, int]:
    """B2 「Comparative Financial Performance - August 2026」에서 (연, 월)."""
    m = re.search(r'-\s*([A-Za-z]+)\s+(20\d{2})', str(ws['B2'].value or ''))
    if not m or m.group(1).upper() not in MONTHS:
        raise ValueError(f'B2 제목에서 월을 읽을 수 없습니다: {ws["B2"].value!r}')
    return int(m.group(2)), MONTHS.index(m.group(1).upper()) + 1


def pnl_blocks(ws, hdr):
    """머리행에서 Cannery..Total 반복 구간을 끊는다. [0]=YTD, [1]=당월, [2]=전년 동월."""
    blocks, cur = [], {}
    for c in range(2, 30):
        lbl = norm(ws.cell(hdr, c).value)
        if lbl in UNITS:
            if lbl in cur:
                blocks.append(cur)
                cur = {}
            cur[lbl] = c
    if cur:
        blocks.append(cur)
    return blocks


def label_at(ws, r) -> str:
    return next((str(ws.cell(r, c).value).strip() for c in (2, 3, 4, 5)
                 if str(ws.cell(r, c).value or '').strip()), '')


def header_row(ws) -> int:
    hdr = next((r + 1 for r in range(1, 12)
                if norm(ws.cell(r, 2).value).upper().startswith('PARTICULAR')), None)
    if hdr is None:
        raise ValueError('PARTICULAR 머리행이 없습니다')
    return hdr


def ytd_by_row(ws) -> dict[int, tuple[str, dict[str, float | None]]]:
    """전월 시트의 YTD 블록 — 행 번호 → (라벨, 부문별 값)."""
    hdr = header_row(ws)
    ytd = pnl_blocks(ws, hdr)[0]
    return {r: (label_at(ws, r), {u: num(ws.cell(r, c).value) for u, c in ytd.items()})
            for r in range(hdr + 1, ws.max_row + 1)}


def prior_ytd_total_col(ws, hdr: int, year: int) -> int | None:
    """머리행 위 「YTD <MONTH> <전년> ACTUAL」 블록의 Total 열. 위치가 판마다 29 → 28 로 움직여 텍스트로 찾는다."""
    for c in range(2, ws.max_column + 1):
        t = norm(ws.cell(hdr - 1, c).value).upper()
        if t.startswith('YTD') and f'{year - 1}ACTUAL' in t:
            for cc in range(c, c + 6):
                if norm(ws.cell(hdr, cc).value) == 'Total':
                    return cc
    return None


def parse_sheet(ws, month: int, prev_ytd=None, year: int = 2026) -> dict:
    hdr = header_row(ws)
    blocks = pnl_blocks(ws, hdr)
    if len(blocks) < 3 or 'Total' not in blocks[1]:
        raise ValueError(f'YTD·당월·전년 블록을 찾지 못했습니다: {blocks}')
    ytd, cur, prev = blocks[0], blocks[1], blocks[2]

    rec: dict = {'month': month}
    lines: dict[str, float] = {}
    lines_prev_ytd: dict[str, float] = {}
    rev_ytd: dict[str, float] = {}
    rev_prev_ytd: dict[str, float] = {}
    prior_col = prior_ytd_total_col(ws, hdr, year)
    printed: dict[str, float | None] = {}
    for r in range(hdr + 1, ws.max_row + 1):
        label = label_at(ws, r)
        if not label:
            continue

        def get(blk, k, r=r, label=label):
            v = num(ws.cell(r, blk[k]).value) if k in blk else None
            if prev_ytd is None or blk is not cur:
                return v
            prev_label, prev_vals = prev_ytd.get(r, ('', {}))
            if prev_label != label:
                raise ValueError(f'{r}행 라벨이 전월과 다릅니다: {prev_label!r} → {label!r}')
            now, before = num(ws.cell(r, ytd[k]).value), prev_vals.get(k)
            if now is None or before is None:
                return v
            diff = now - before
            if abs((v or 0) - diff) > 0.5:
                printed[f'{r}:{label}:{k}'] = rnd(v)
                return diff
            return v

        key = PNL_ROWS.get(norm(label).lower())
        if key and key not in rec:
            rec[key] = rnd(get(cur, 'Total'))
            rec[key + 'Ytd'] = rnd(get(ytd, 'Total'))
            rec[key + 'Prev'] = rnd(get(prev, 'Total'))
            for u in ('Cannery', 'Fishmeal', 'CBU', 'FBU'):
                rec[f'{key}_{u.lower()}'] = rnd(get(cur, u))
        if label in REVENUE_LINES and label not in rev_ytd and 'revenue' in rec and 'cos' not in rec:
            rev_ytd[label] = rnd(num(ws.cell(r, ytd['Total']).value) or 0)
            if prior_col:
                rev_prev_ytd[label] = rnd(num(ws.cell(r, prior_col).value) or 0)
        if label in COST_LINES and label not in lines:
            v = get(cur, 'Total')
            if v is not None:
                lines[label] = rnd(v)
            pv = num(ws.cell(r, prior_col).value) if prior_col else None
            if pv is not None:
                lines_prev_ytd[label] = rnd(pv)
    rec['costLines'] = lines
    # 전년 동기 누계(원가 계정) — 「전기료 YTD 전년 대비」 같은 문장을 손으로 적지 않게 한다
    rec['costLinesPrevYtd'] = lines_prev_ytd
    # 매출 세부 계정 누계(당해·전년) — 「전년 대비 매출 갭 중 사업 축소분」을 손으로 적지 않게 한다
    rec['revenueLinesYtd'] = rev_ytd
    rec['revenueLinesPrevYtd'] = rev_prev_ytd
    if printed:
        rec['correction'] = {
            'method': '당월 = 이번 YTD − 전월 YTD (인쇄 당월값이 YTD 차분과 어긋난 칸만)',
            'printed': printed,
        }
    for r in range(1, 8):
        for c in range(8, 18):
            lbl = norm(ws.cell(r, c).value)
            if lbl == 'Skipjack' and rec.get('fishPriceSJ') is None:
                rec['fishPriceSJ'] = rnd(num(ws.cell(r, c + 1).value))
            elif lbl == 'Yellowfin' and rec.get('fishPriceYF') is None:
                rec['fishPriceYF'] = rnd(num(ws.cell(r, c + 1).value))
            elif lbl.startswith('$1=') and rec.get('forex') is None:
                rec['forex'] = rnd(num(lbl.split('=')[-1]), 4)

    missing = [k for k in ('revenue', 'cos', 'gp', 'sga', 'op', 'nonop', 'interest', 'net')
               if rec.get(k) is None]
    if missing:
        raise ValueError(f'당월 합계가 비어 있습니다: {missing}')
    return rec


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('workbook', type=Path)
    ap.add_argument('--sheet', help='시트 이름(기본: 첫 시트)')
    ap.add_argument('--prev-workbook', type=Path, help='전월 시트가 든 파일(YTD 대조용)')
    ap.add_argument('--prev-sheet', help='전월 시트 이름')
    ap.add_argument('--data', type=Path, default=Path('public/data/cosmo/cosmo_2026.json'))
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]
    year, month = title_month(ws)
    if year != 2026:
        raise ValueError(f'{year}년 손익은 cosmo_2026.json 대상이 아닙니다')
    prev_ytd = None
    if args.prev_workbook:
        pwb = openpyxl.load_workbook(args.prev_workbook, data_only=True)
        pws = pwb[args.prev_sheet] if args.prev_sheet else pwb.worksheets[0]
        if title_month(pws) != (year, month - 1):
            raise ValueError(f'전월 시트가 {month - 1}월이 아닙니다: {title_month(pws)}')
        prev_ytd = ytd_by_row(pws)
    rec = parse_sheet(ws, month, prev_ytd, year)
    rec['source'] = unicodedata.normalize('NFC', args.workbook.name)
    rec['sha256'] = hashlib.sha256(args.workbook.read_bytes()).hexdigest()

    payload = json.loads(args.data.read_text(encoding='utf-8'))
    before = next((m for m in payload['monthly'] if m['month'] == month - 1), None)
    if before:  # 계정 합계의 YTD 연속성 — 어긋나면 쓰지 않는다
        gaps = {k: round(before[k + 'Ytd'] + rec[k] - rec[k + 'Ytd'], 2)
                for k in ('revenue', 'cos', 'gp', 'sga', 'op', 'nonop', 'interest', 'net')
                if abs(before[k + 'Ytd'] + rec[k] - rec[k + 'Ytd']) > 1}
        if gaps:
            raise ValueError(f'전월 YTD + 당월 ≠ 이번 YTD: {gaps} — --prev-workbook 로 행 단위 대조')
    monthly = [m for m in payload['monthly'] if m['month'] != month] + [rec]
    monthly.sort(key=lambda m: m['month'])
    if [m['month'] for m in monthly] != list(range(1, len(monthly) + 1)):
        raise ValueError(f'월 계열에 빈 달이 생깁니다: {[m["month"] for m in monthly]}')
    payload['monthly'] = monthly
    payload['meta']['monthCount'] = len(monthly)

    print(json.dumps({k: rec[k] for k in ('month', 'revenue', 'gp', 'op', 'net', 'netYtd',
                                          'fishPriceSJ', 'forex', 'sha256')},
                     ensure_ascii=False, indent=2))
    if not args.dry_run:
        args.data.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')),
                             encoding='utf-8')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
